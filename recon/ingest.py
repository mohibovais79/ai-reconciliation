"""Read CSV/XLSX, find the header row, drop junk rows (with reasons), auto-map columns."""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz import fuzz

from recon.config import Settings
from recon.models import DroppedRow, Flag, FlagType, SourceRow
from recon.normalize import (
    canonical_unit,
    common_prefix,
    detect_dayfirst,
    is_ambiguous_date,
    normalize_item,
    normalize_key,
    normalize_lot,
    normalize_supplier,
    parse_date,
    parse_datetime,
    parse_qty,
    to_base,
    year_from_label,
)

CANONICAL_FIELDS = ["key", "date", "supplier", "item", "lot", "qty", "unit",
                    "entered_by", "entered_at", "remarks"]
TOTAL_RE = re.compile(r"^\s*(grand\s*)?(sub\s*)?total\b|^\s*sum\b|^\s*totals?\s*:?\s*$", re.I)


class IngestError(ValueError):
    """User-facing problem reading a file."""


@dataclass
class Table:
    """A source file after header detection and junk-row removal."""

    source: str
    name: str
    sha256: str
    columns: list[str]
    records: list[tuple[str, dict]]          # (row_ref like 'r17', {column: value})
    dropped: list[DroppedRow] = field(default_factory=list)
    rows_read: int = 0                       # every physical row except the header
    sheet: str | None = None
    header_row: int = 1
    meta: dict = field(default_factory=dict)  # register: per-row cell info, page labels, etc.

    @property
    def fingerprint(self) -> str:
        """Stable hash of the header set, used to recall saved mappings."""
        norm = sorted(_norm_header(c) for c in self.columns)
        return hashlib.sha256(f"{self.source}|{'|'.join(norm)}".encode()).hexdigest()[:16]


def _norm_header(h: object) -> str:
    s = re.sub(r"\(.*?\)", " ", str(h or "")).casefold()
    s = s.replace("#", " no ").replace("_", " ")
    return " ".join(re.sub(r"[^\w\s]", " ", s).split())


def _blank(v: object) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def list_sheets(data: bytes) -> list[str]:
    """Sheet names of an XLSX file."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read_csv(data: bytes) -> list[list[object]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes, sheet: str | None) -> tuple[list[list[object]], str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        return rows, ws.title
    finally:
        wb.close()


def _is_texty(v: object) -> bool:
    if _blank(v) or not isinstance(v, str):
        return False
    return parse_qty(v)[0] is None and parse_date(v, False) is None


def detect_header(rows: list[list[object]]) -> int:
    """Index of the first row that looks like a header (mostly non-numeric text, wide enough)."""
    width = max((sum(not _blank(v) for v in r) for r in rows[:50]), default=0)
    for i, r in enumerate(rows[:50]):
        filled = [v for v in r if not _blank(v)]
        if len(filled) < max(2, width * 0.5):
            continue
        if sum(_is_texty(v) for v in filled) >= 0.7 * len(filled):
            return i
    raise IngestError("Could not find a header row (a row of column names).")


def read_table(source: str, name: str, data: bytes, settings: Settings,
               sheet: str | None = None) -> Table:
    """Read a CSV/XLSX upload into a Table, recording every dropped row and why."""
    if len(data) > settings.app.max_file_mb * 1024 * 1024:
        raise IngestError(f"{name}: larger than {settings.app.max_file_mb} MB.")
    ext = Path(name).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        if not data.startswith(b"PK"):
            raise IngestError(f"{name}: not a valid Excel file.")
        rows, sheet = _read_xlsx(data, sheet)
    elif ext in (".csv", ".txt", ".tsv"):
        rows = _read_csv(data)
    else:
        raise IngestError(f"{name}: unsupported file type (use CSV or XLSX).")
    if not rows:
        raise IngestError(f"{name}: file is empty.")
    h = detect_header(rows)
    raw_cols = rows[h]
    columns, seen = [], {}
    for i, c in enumerate(raw_cols):
        base = str(c).strip() if not _blank(c) else f"Column{i + 1}"
        seen[base] = seen.get(base, 0) + 1
        columns.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    header_norm = [_norm_header(c) for c in raw_cols]
    t = Table(source=source, name=name, sha256=hashlib.sha256(data).hexdigest(),
              columns=columns, records=[], sheet=sheet, header_row=h + 1)
    for idx, r in enumerate(rows):
        if idx == h:
            continue
        t.rows_read += 1
        ref = f"r{idx + 1}"
        rec = {c: (r[j] if j < len(r) else None) for j, c in enumerate(columns)}
        rec = {k: (v.strip() if isinstance(v, str) else v) for k, v in rec.items()}
        raw = {k: _jsonable(v) for k, v in rec.items() if not _blank(v)}
        if idx < h:
            reason = "above header (title/notes)"
        elif all(_blank(v) for v in rec.values()):
            reason = "blank row"
        elif [_norm_header(v) for v in r[: len(raw_cols)]] == header_norm:
            reason = "repeated header row"
        elif any(isinstance(v, str) and TOTAL_RE.match(v) for v in rec.values()):
            reason = "totals row"
        else:
            reason = None
        if reason:
            t.dropped.append(DroppedRow(source=source, row_ref=ref, reason=reason, raw=raw))
        else:
            t.records.append((ref, rec))
    return t


def _jsonable(v: object) -> object:
    if isinstance(v, dt.datetime | dt.date):
        return v.isoformat()
    return v


# ---------- auto column mapping ----------

def _header_score(col: str, synonyms: list[str]) -> float:
    h = _norm_header(col)
    if not h:
        return 0.0
    return max(max(fuzz.ratio(h, s), fuzz.token_sort_ratio(h, s)) for s in synonyms)


def _value_fraction(values: list[object], fieldname: str, settings: Settings) -> float | None:
    """Share of sample values that look right for a field (None = no check)."""
    vals = [v for v in values if not _blank(v)]
    if not vals:
        return None
    if fieldname in ("date", "entered_at"):
        ok = sum(isinstance(v, dt.date) or parse_date(v, False) is not None
                 or parse_date(v, True) is not None for v in vals
                 if not (isinstance(v, int | float) and not 20000 < float(v) < 80000))
        return ok / len(vals)
    if fieldname == "qty":
        return sum(parse_qty(v)[0] is not None for v in vals) / len(vals)
    if fieldname == "unit":
        return sum(canonical_unit(v, settings) is not None for v in vals) / len(vals)
    return None


def unit_hint_from_header(col: str, settings: Settings) -> str | None:
    """Unit written in a header like 'Qty (yds)'."""
    m = re.search(r"\(([^)]+)\)", col)
    return canonical_unit(m.group(1), settings) if m else None


def auto_map(table: Table, settings: Settings) -> dict[str, str | None]:
    """Propose {canonical_field: column} using synonyms, fuzzy headers and value checks."""
    samples = {c: [rec.get(c) for _, rec in table.records[:40]] for c in table.columns}
    scored = []
    for f in CANONICAL_FIELDS:
        syn = settings.mapping.synonyms.get(f, [f])
        for c in table.columns:
            s = _header_score(c, syn)
            frac = _value_fraction(samples[c], f, settings)
            if frac is not None:
                s += 10 if frac >= 0.8 else (-40 if frac < 0.3 else 0)
            scored.append((s, f, c))
    scored.sort(key=lambda x: -x[0])
    mapping: dict[str, str | None] = {f: None for f in CANONICAL_FIELDS}
    used: set[str] = set()
    for s, f, c in scored:
        if s < settings.mapping.header_min_score or mapping[f] or c in used:
            continue
        mapping[f] = c
        used.add(c)
    return mapping


def missing_required(mapping: dict[str, str | None], settings: Settings) -> list[str]:
    """Required canonical fields that have no column."""
    req = list(settings.mapping.required)
    if not mapping.get("key"):
        req += [f for f in ("date", "supplier", "item") if f not in req]
        req = [f for f in req if f != "key"]
    return [f for f in req if not mapping.get(f)]


# ---------- register pages -> Table ----------

def _align_to(base: list[str], cols: list[str]) -> dict[str, str]:
    """Map a page's headers onto the first page's headers."""
    if len(cols) == len(base):
        return dict(zip(cols, base, strict=True))
    return {c: max(base, key=lambda b: fuzz.ratio(_norm_header(b), _norm_header(c))) for c in cols}


def register_table(results: list, name: str = "register pages") -> Table:
    """Turn successful page extractions into one Table with per-cell metadata."""
    ok = [r for r in results if r.status == "ok" and r.extraction]
    base_cols = ok[0].extraction.columns if ok else []
    sha = hashlib.sha256("".join(r.sha256 for r in results).encode()).hexdigest()
    t = Table(source="REG", name=name, sha256=sha, columns=list(base_cols), records=[],
              meta={"cells": {}, "row_flags": {}, "page_flags": [], "periods": [],
                    "page_of": {}})
    for r in ok:
        ext = r.extraction
        colmap = _align_to(base_cols, ext.columns)
        if ext.period_label:
            t.meta["periods"].append(ext.period_label)
        p = r.page_index + 1
        for f in r.page_flags:
            t.meta["page_flags"].append((p, f))
        if ext.rows_visible is not None and ext.rows_visible != len(ext.rows):
            t.meta["page_flags"].append((p, Flag(
                type=FlagType.ROW_COUNT_MISMATCH,
                message=f"page shows {ext.rows_visible} filled rows, {len(ext.rows)} read")))
        for i, row in enumerate(ext.rows):
            ref = f"p{p}r{i + 1}"
            t.rows_read += 1
            cells = {colmap[k]: v for k, v in row.items()}
            rec = {c: (cells[c].text if c in cells else None) for c in base_cols}
            if all(_blank(v) and not (cells.get(c) and cells[c].struck_text)
                   for c, v in rec.items()):
                t.dropped.append(DroppedRow(source="REG", row_ref=ref, reason="blank row"))
                continue
            t.records.append((ref, rec))
            t.meta["cells"][ref] = cells
            t.meta["page_of"][ref] = p
            flags = []
            for f in r.flags_by_row.get(i, []):
                f2 = f.model_copy()
                if f2.field:
                    f2.field = colmap.get(f2.field, f2.field)
                flags.append(f2)
            t.meta["row_flags"][ref] = flags
    return t


# ---------- mapped Table -> canonical SourceRows ----------

@dataclass
class SourceContext:
    """Per-source parsing decisions (inferred once per column)."""

    source: str
    mapping: dict[str, str | None]
    key_prefix: str = ""
    dayfirst: bool = False
    dayfirst_known: bool = True
    default_year: int | None = None
    has_real_dates: bool = False
    unit_hint: str | None = None


def build_context(table: Table, mapping: dict[str, str | None], settings: Settings,
                  period_year: int | None = None) -> SourceContext:
    """Infer key prefix, date order, year and header unit hint for one source."""
    def col(f):
        c = mapping.get(f)
        return [rec.get(c) for _, rec in table.records] if c else []
    keys = [str(v) for v in col("key") if not _blank(v)]
    dates = col("date")
    df = detect_dayfirst(dates)
    year = period_year
    if year is None and table.meta.get("periods"):
        year = next((y for y in map(year_from_label, table.meta["periods"]) if y), None)
    qty_col = mapping.get("qty")
    return SourceContext(
        source=table.source, mapping=mapping, key_prefix=common_prefix(keys),
        dayfirst=settings.dates.default_dayfirst if df is None else df,
        dayfirst_known=df is not None or not any(is_ambiguous_date(v) for v in dates),
        default_year=year,
        has_real_dates=any(isinstance(v, dt.date) for v in dates),
        unit_hint=unit_hint_from_header(qty_col, settings) if qty_col else None)


def set_field(row: SourceRow, fieldname: str, value: object, ctx: SourceContext,
              settings: Settings, learned: dict[str, dict[str, str]] | None = None) -> None:
    """Normalise one raw value onto a SourceRow (also used when a reviewer enters a value)."""
    learned = learned or {}
    blank = _blank(value)
    if fieldname == "key":
        row.key_raw = None if blank else str(value)
        row.key = None if blank else normalize_key(value, ctx.key_prefix)
    elif fieldname == "date":
        row.date_raw = None if blank else str(value)
        row.date = None if blank else parse_date(value, ctx.dayfirst, ctx.default_year)
        row.date_was_text = isinstance(value, str) and ctx.has_real_dates and not blank
        if not blank and row.date is None:
            row.flags.append(Flag(type=FlagType.BAD_FORMAT, field="date", blocking=False,
                                  message=f"date '{value}' could not be read"))
        elif not blank and not ctx.dayfirst_known and is_ambiguous_date(value):
            row.flags.append(Flag(type=FlagType.AMBIGUOUS_DATE, field="date", blocking=False,
                                  message=f"'{value}' could be day/month or month/day"))
    elif fieldname == "supplier":
        row.supplier_raw = None if blank else str(value)
        norm = normalize_supplier(value, settings)
        row.supplier_norm = learned.get("supplier", {}).get(norm, norm) if norm else None
    elif fieldname == "item":
        row.item_raw = None if blank else str(value)
        row.item_norm = normalize_item(value)
    elif fieldname == "lot":
        row.lot = normalize_lot(value)
        row.extras["lot_raw"] = None if blank else str(value)
    elif fieldname == "qty":
        row.qty_raw = None if blank else (value if isinstance(value, int | float) else str(value))
        q, suffix = parse_qty(value)
        row.qty = q
        if suffix and not row.unit_raw:
            row.extras["qty_unit_suffix"] = suffix
        if not blank and q is None:
            row.flags.append(Flag(type=FlagType.BAD_FORMAT, field="qty",
                                  message=f"quantity '{value}' is not a number"))
    elif fieldname == "unit":
        row.unit_raw = None if blank else str(value)
    elif fieldname == "entered_at":
        row.entered_at = parse_datetime(value, ctx.dayfirst)
        row.extras["entered_at"] = None if blank else str(value)
    else:
        row.extras[fieldname] = None if blank else value


def finish_units(row: SourceRow, ctx: SourceContext, settings: Settings,
                 learned: dict[str, dict[str, str]] | None = None) -> None:
    """Resolve the canonical unit and base quantity after all fields are set."""
    raw_u = row.unit_raw or row.extras.get("qty_unit_suffix")
    unit = canonical_unit(raw_u, settings, (learned or {}).get("unit")) if raw_u else ctx.unit_hint
    row.flags = [f for f in row.flags if f.type != FlagType.UNKNOWN_UNIT]
    if raw_u and unit is None:
        row.flags.append(Flag(type=FlagType.UNKNOWN_UNIT, field="unit", blocking=False,
                              message=f"unit '{raw_u}' is not in the unit list"))
    row.unit = unit
    row.qty_base, row.base_unit = to_base(row.qty, unit, settings)


def _cell_flags(ref: str, table: Table, colfield: dict[str, str]) -> list[Flag]:
    """Doubts from the register reading (unclear, struck, two-pass disagreement)."""
    flags = []
    for col, cell in table.meta.get("cells", {}).get(ref, {}).items():
        f = colfield.get(col)
        if not f:
            continue
        if cell.legibility != "clear":
            flags.append(Flag(type=FlagType.UNCLEAR_FIELD, field=f,
                              message=f"{cell.legibility}: read '{cell.text or ''}'"
                                      + (f" ({cell.note})" if cell.note else "")))
        if cell.struck_text:
            flags.append(Flag(type=FlagType.STRUCK_VALUE, field=f, suggestion=cell.text,
                              candidates=[cell.text or "", cell.struck_text],
                              message=f"crossed-out {cell.struck_text}, written {cell.text}"))
    for fl in table.meta.get("row_flags", {}).get(ref, []):
        if fl.field is None or fl.field in colfield:
            f2 = fl.model_copy()
            f2.field = colfield.get(fl.field) if fl.field else None
            flags.append(f2)
    return flags


def to_source_rows(table: Table, mapping: dict[str, str | None], settings: Settings,
                   period_year: int | None = None,
                   learned: dict[str, dict[str, str]] | None = None
                   ) -> tuple[list[SourceRow], SourceContext]:
    """Build canonical SourceRows for every ingested record of a table."""
    ctx = build_context(table, mapping, settings, period_year)
    colfield = {c: f for f, c in mapping.items() if c}
    rows = []
    for ref, rec in table.records:
        row = SourceRow(source=table.source, row_id=f"{table.source}:{ref}",
                        raw={k: _jsonable(v) for k, v in rec.items() if not _blank(v)})
        for f, c in mapping.items():
            if c:
                set_field(row, f, rec.get(c), ctx, settings, learned)
        finish_units(row, ctx, settings, learned)
        if table.source == "REG":
            row.extras["page"] = table.meta["page_of"].get(ref)
            cf = _cell_flags(ref, table, colfield)
            if any(f.field == "qty" and f.type == FlagType.UNCLEAR_FIELD for f in cf):
                row.qty, row.qty_base = None, None      # never use a guessed quantity
            row.flags.extend(cf)
        rows.append(row)
    return rows, ctx
