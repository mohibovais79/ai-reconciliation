"""M8: workbook tabs, live-formula proof, injection sanitising, summary and number check."""

import io

from openpyxl import load_workbook

from recon.report import (
    build_workbook,
    ledger_csv,
    mismatched_csv,
    sanitize,
    summary_text,
    verify_numbers,
)
from recon.store import Store
from tests.test_pipeline_review import S, _confirm_all, _run


def _confirmed_run():
    store = Store(":memory:")
    _confirm_all(store, _run(store))
    return _run(store), store


def test_workbook_tabs_and_formulas():
    run, store = _confirmed_run()
    wb = load_workbook(io.BytesIO(build_workbook(run, S, store.all_decisions())))
    assert wb.sheetnames == ["Summary", "Action lists", "Mismatched", "Missing",
                             "Duplicates & non-receipts", "Matched", "Review log", "Ledger",
                             "Proof", "Inputs"]
    assert wb["Summary"]["A1"].value.startswith("RUN STATUS: COMPLETE")
    assert wb["Ledger"].max_row - 1 == len(run.ledger)
    formulas = [c.value for row in wb["Proof"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("COUNTIFS(Ledger!" in f for f in formulas)
    assert any("SUMIFS(Ledger!" in f for f in formulas)
    assert wb["Review log"].max_row >= 2


def test_sanitize_and_csv():
    assert sanitize("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert sanitize("@SUM") == "'@SUM" and sanitize("abc") == "abc" and sanitize(5) == 5
    run, _ = _confirmed_run()
    assert ledger_csv(run).decode("utf-8-sig").count("\n") == len(run.ledger) + 1
    assert "qty" in mismatched_csv(run, S).decode("utf-8-sig")


def test_summary_short_exceptions_first_and_verified():
    run, _ = _confirmed_run()
    text = summary_text(run, S, "recon.xlsx")
    lines = text.splitlines()
    assert len(lines) <= 14
    assert lines[1].startswith("NEEDS ACTION")
    assert lines[-2].startswith("PROOF: all") and "✓" in lines[-2]
    assert verify_numbers(text, run, S)
    assert not verify_numbers(text + " 987654", run, S)
